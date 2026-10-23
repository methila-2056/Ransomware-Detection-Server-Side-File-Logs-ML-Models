"""Build an IEEE-conference-style DOCX from the markdown draft.
Produces conference/conference_paper.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).resolve().parent
FIG = BASE / "figures"
SRC = BASE / "conference_paper_draft.md"
OUT = BASE / "conference_paper.docx"

FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)")


def set_cols(section, num=2, space_twips=240):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


MATH = re.compile(r"((?:[Ccrdxy])_(?:TP|TN|FP|FN|i))")


def _emit(p, text, size=None, italic=False, bold=False, code=False):
    if code:
        r = p.add_run(text)
        r.font.name = "Courier New"
        if size:
            r.font.size = Pt(size)
        return
    for part in MATH.split(text):
        if not part:
            continue
        m = re.fullmatch(r"([Ccrdxy])_(TP|TN|FP|FN|i)", part)
        if m:
            for sub_text, sub in ((m.group(1), False), (m.group(2), True)):
                r = p.add_run(sub_text)
                r.italic = italic
                r.bold = bold
                if sub:
                    r.font.subscript = True
                if size:
                    r.font.size = Pt(size)
        else:
            r = p.add_run(part)
            r.italic = italic
            r.bold = bold
            if size:
                r.font.size = Pt(size)


def add_runs(p, text, size=None, base_italic=False, base_bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _emit(p, part[2:-2], size, italic=base_italic, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _emit(p, part[1:-1], size, italic=True, bold=base_bold)
        elif part.startswith("`") and part.endswith("`"):
            _emit(p, part[1:-1], size, code=True)
        else:
            _emit(p, part, size, italic=base_italic, bold=base_bold)
    return p


def para(doc, text, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False,
         bold=False, space_after=0, first_indent=0.2, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0
    if first_indent:
        pf.first_line_indent = Inches(first_indent)
    add_runs(p, text, size=size, base_italic=italic, base_bold=bold)
    if keep:
        pf.keep_with_next = True
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(10 if level == 1 else 6)
    pf.space_after = Pt(4)
    pf.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(10)
    if level == 1:
        r.font.small_caps = True
        r.bold = False
    else:
        r.italic = True
    return p


def add_image(doc, path, width=3.2):
    if not path.exists():
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    return True


def add_table(doc, rows):
    if not rows:
        return None
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            val = row[j] if j < len(row) else ""
            add_runs(p, val, size=8)
            for r in p.runs:
                if i == 0:
                    r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def parse_md():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    return lines


def build():
    doc = Document()

    # page setup
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(0.625)
    sec.right_margin = Inches(0.625)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    lines = parse_md()
    i = 0
    # --- Title / author block (single column) ---
    while i < len(lines) and not lines[i].startswith("# "):
        i += 1
    title = lines[i][2:].strip() if i < len(lines) else "Paper"
    i += 1

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.font.name = FONT; r.font.size = Pt(20); r.bold = True
    p.paragraph_format.space_after = Pt(8)

    # author lines until '---'
    author_lines = []
    while i < len(lines) and lines[i].strip() != "---":
        s = lines[i].strip()
        if s and not s.startswith(">"):
            author_lines.append(s)
        i += 1
    for a in author_lines:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, a, size=11)

    # author placeholder note (blockquote) as italic small
    j = i
    note = []
    while j < len(lines) and lines[j].strip() != "---":
        if lines[j].strip().startswith(">"):
            note.append(lines[j].strip("> ").strip())
        j += 1
    for n in note:
        para(doc, n, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True,
             first_indent=0, space_after=2)

    # --- switch to two columns for the body ---
    new_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    new_sec.top_margin = Inches(0.75)
    new_sec.bottom_margin = Inches(1.0)
    new_sec.left_margin = Inches(0.625)
    new_sec.right_margin = Inches(0.625)
    set_cols(new_sec, 2, 240)

    # find start of body: first '**Abstract**' or section heading
    body_start = None
    for k in range(i, len(lines)):
        if lines[k].startswith("**Abstract**") or lines[k].startswith("## "):
            body_start = k
            break
    if body_start is None:
        body_start = i

    k = body_start
    in_refs = False
    pending_table_title = None
    while k < len(lines):
        raw = lines[k]
        s = raw.strip()

        if not s or s == "---":
            k += 1
            continue

        if s.startswith("> "):
            # leftover blockquote note -> italic small
            para(doc, s[2:].strip(), size=8, italic=True, first_indent=0)
            k += 1
            continue

        if s.startswith("## "):
            in_refs = s[2:].strip().upper().startswith("REFERENCES")
            heading(doc, s[3:].strip(), level=1)
            k += 1
            continue

        if s.startswith("### "):
            heading(doc, s[4:].strip(), level=2)
            k += 1
            continue

        if s in ("!EQ1!", "!EQ2!"):
            from docx.enum.text import WD_TAB_ALIGNMENT
            num = "1" if s == "!EQ1!" else "2"
            img = FIG / ("eq1.png" if num == "1" else "eq2.png")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.space_before = Pt(4)
            pf.space_after = Pt(4)
            pf.tab_stops.add_tab_stop(Inches(3.35), WD_TAB_ALIGNMENT.RIGHT)
            if img.exists():
                p.add_run().add_picture(str(img), width=Inches(1.7 if num == "1" else 3.1))
            r = p.add_run("\t(" + num + ")")
            r.font.size = Pt(10)
            k += 1
            continue

        if s.startswith("**TABLE") or s.startswith("**Table"):
            # table title line (may include trailing text after em dash)
            pending_table_title = s.strip("*")
            k += 1
            # skip blank lines between title and table
            while k < len(lines) and not lines[k].strip():
                k += 1
            # collect markdown table rows
            rows = []
            while k < len(lines) and lines[k].strip().startswith("|"):
                cells = [c.strip() for c in lines[k].strip().strip("|").split("|")]
                if all(set(c) <= set("-: ") for c in cells):
                    k += 1
                    continue
                rows.append(cells)
                k += 1
            # title paragraph (small caps, centered) above table
            tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tp.paragraph_format.space_before = Pt(6)
            tp.paragraph_format.space_after = Pt(2)
            tr = tp.add_run(pending_table_title); tr.font.size = Pt(8)
            tr.font.small_caps = True
            add_table(doc, rows)
            continue
        if s.startswith("**Figure") or s.startswith("**Fig."):
            # extract figure path if present
            path = None
            m = re.search(r"\(figure:\s*([^)]+)\)", s)
            if m:
                path = FIG.parent / m.group(1).strip()
                s = re.sub(r"\(figure:[^)]+\)", "", s).strip()
            if path:
                add_image(doc, path, width=3.25)
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(6)
            add_runs(cp, s.replace("**", "").strip(), size=8)
            k += 1
            continue

        if in_refs and re.match(r"^\[\d+\]", s):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            add_runs(p, s, size=8)
            k += 1
            continue

        # normal paragraph
        para(doc, s, size=10)
        k += 1

    doc.save(str(OUT))
    print("Saved", OUT, OUT.stat().st_size, "bytes")


if __name__ == "__main__":
    build()
